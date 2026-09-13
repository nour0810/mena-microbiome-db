#!/usr/bin/env python
"""Step 12. Rebuild Supplementary Tables S1-S3 on the v1.2 corpus and update S4's summary."""
import json, numpy as np, pandas as pd
from openpyxl import load_workbook, Workbook

import os
from pathlib import Path
W = str(Path(os.environ.get('MENA_OUT', Path(__file__).resolve().parent / 'supplementary_out')))
os.makedirs(W, exist_ok=True)
d = pd.read_csv('corpus_v12.tsv', sep='\t', dtype=str, low_memory=False)
d['rc'] = pd.to_numeric(d.read_count, errors='coerce')
d['bc'] = pd.to_numeric(d.base_count, errors='coerce')
d['yr'] = pd.to_numeric(d.year, errors='coerce')
N = len(d)
R = json.load(open('run_v12/data/mena_rigorous_results.json'))
PC = json.load(open('percategory_v12.json'))
BS = json.load(open('biosample_both.json'))['v12']
F11 = json.load(open('fig11_values.json')); F12 = json.load(open('fig12_values.json'))
F10 = json.load(open('fig10_values.json'))
PH = {'not collected','not applicable','not available','unknown','missing','none','null','na','n/a','-',''}
def filled(s): return ~s.fillna('').astype(str).str.strip().str.lower().isin(PH)
NOTE = ("MENA Microbiome Database v1.2 (56,494 runs). Supersedes v1.1 (60,126 runs): 4,423 runs in "
        "studies the BioProject-level census judged not to be community metagenomes were removed, and "
        "791 community marker-gene runs recovered from the rejected arm were added.")

def sheet(wb, name, df, note=None):
    ws = wb.create_sheet(name[:31])
    ws.append(list(df.columns))
    for _, r in df.iterrows(): ws.append([None if pd.isna(v) else v for v in r.tolist()])
    if note: ws.append([]); ws.append([note])

# ------------------------------------------------------------------- S1
wb = Workbook(); wb.remove(wb.active)
sheet(wb, 'headline_stats', pd.DataFrame({
 'Metric': ['Total metagenomic runs','Unique samples','Unique BioProjects','MENA countries represented',
            'Year range','Shotgun metagenomics runs','Amplicon metagenomics runs','Metatranscriptomics runs',
            'Human-associated runs','Environmental runs','Animal runs','Plant runs'],
 'Value': [N, d.sample_accession.nunique(), d.bioproject.nunique(), d.country.nunique(), '2013-2026',
           int((d.data_subtype=='shotgun_metagenomics').sum()),
           int((d.data_subtype=='amplicon_metagenomics').sum()),
           int((d.data_subtype=='metatranscriptomics').sum()),
           int((d.broad_category=='Human').sum()), int((d.broad_category=='Environment').sum()),
           int((d.broad_category=='Animal').sum()), int((d.broad_category=='Plant').sum())]}), NOTE)
g = d.groupby('country').agg(n_runs=('run_accession','size'), n_bioprojects=('bioproject','nunique'),
                             n_samples=('sample_accession','nunique')).sort_values('n_runs', ascending=False)
sheet(wb, '01_geographic_distribution', g.reset_index(),
      'The per-country n_bioprojects column is not additive; use headline_stats for the corpus total.')
sheet(wb, '02_temporal_trends', d.yr.value_counts().sort_index().rename_axis('year')
      .reset_index(name='n_runs'))
bc = d.broad_category.value_counts().rename_axis('broad_category').reset_index(name='count')
sheet(wb, '03a_broad_categories', bc,
      'Assignment accuracy for these categories is quantified in Supplementary Table S6 and reported in §3.1.')
hs = (d[d.broad_category.isin(['Human','Animal'])].groupby(['broad_category','specific_category'])
      .size().reset_index(name='n_runs').sort_values('n_runs', ascending=False))
hs.columns = ['broad','specific','n_runs']; sheet(wb, '04_host distribution', hs)
sheet(wb, '05_sequencing_platforms', d.instrument_platform.value_counts()
      .rename_axis('instrument_platform').reset_index(name='count'))
sheet(wb, '06a_library_strategy', d.library_strategy.value_counts()
      .rename_axis('library_strategy').reset_index(name='count'))
sheet(wb, '06b_library_source', d.library_source.value_counts()
      .rename_axis('library_source').reset_index(name='count'))
sp = pd.DataFrame(BS['top_overall'][:30])
sp['% of corpus'] = (100*sp.n_runs/N).round(2); sheet(wb, '03b_specific_categories', sp)
old = load_workbook(f'{W}/Supplementary_Table_S1.xlsx', data_only=True)
if '07_excluded_non_MENA' in old.sheetnames:
    ws = wb.create_sheet('07_excluded_non_MENA')
    for row in old['07_excluded_non_MENA'].iter_rows(values_only=True): ws.append(list(row))
wb.save(f'{W}/Supplementary_Table_S1.xlsx'); print('S1 rebuilt')

# ------------------------------------------------------------------- S2
wb = Workbook(); wb.remove(wb.active)
im = d.instrument_model.value_counts().head(15).rename_axis('Instrument Model').reset_index(name='Runs (exact)')
im['% of Total'] = (100*im['Runs (exact)']/N).round(2); sheet(wb, 'Instrument Models', im, NOTE)
MIXS = [('Geographic location','country'),('Collection date','collection_date'),
        ('GPS coordinates','sample_lat_lon'),('Environmental biome','environment_biome'),
        ('Environmental feature','environment_feature'),('Environmental material','environment_material'),
        ('Isolation source','isolation_source_meta'),('Host (if applicable)','host_meta'),
        ('Library strategy','library_strategy'),('Sequencing platform','instrument_platform'),
        ('Project name','bioproject'),('Sample title','sample_title')]
rows = []
for c in d.country.value_counts().index:
    s = d[d.country == c]; rec = {'country': c, 'n_runs': len(s)}
    for lab, col in MIXS:
        cc = col if col in d.columns else col.replace('_meta','')
        rec[lab] = round(100*filled(s[cc]).mean(), 2)
    rows.append(rec)
mc = pd.DataFrame(rows); mc['mean'] = mc[[l for l,_ in MIXS]].mean(axis=1).round(2)
sheet(wb, 'metadata completeness', mc,
      f"Corpus mean {R['mixs']['overall_mean_pct']}%, median per run {R['mixs']['overall_median_pct']}%.")
lp = pd.concat([
  d.library_layout.fillna('unspecified').str.upper().value_counts().rename_axis('Subcategory')
    .reset_index(name='Run Count').assign(Category='Library Layout'),
  d.library_selection.fillna('unspecified').value_counts().rename_axis('Subcategory')
    .reset_index(name='Run Count').assign(Category='Library Selection'),
  d.library_strategy.value_counts().rename_axis('Subcategory')
    .reset_index(name='Run Count').assign(Category='Library Strategy')])
lp['% of Total'] = (100*lp['Run Count']/N).round(2)
sheet(wb, 'library prep', lp[['Category','Subcategory','Run Count','% of Total']],
      'Recomputed directly from the released library_layout and library_selection fields. These values '
      'supersede those in the first submission, which did not reproduce from the deposit (§3.5).')
amp = d[d.data_subtype == 'amplicon_metagenomics']
tg = amp.sample_target_gene.fillna('').astype(str).str.strip(); ok = filled(amp.sample_target_gene)
def norm(v):
    u = v.upper()
    return '16S rRNA' if '16S' in u else 'ITS (fungal)' if 'ITS' in u else '18S rRNA' if '18S' in u else 'other'
mk = tg[ok].map(norm).value_counts().rename_axis('Marker').reset_index(name='Runs')
mk['% of annotated subset'] = (100*mk.Runs/mk.Runs.sum()).round(1)
sheet(wb, 'target genes', mk,
      f"Structured marker annotation covers only {int(ok.sum()):,} of {len(amp):,} amplicon runs "
      f"({100*ok.mean():.1f}%); target_subfragment {int(filled(amp.sample_target_subfragment).sum()):,}, "
      f"pcr_primers {int(filled(amp.sample_pcr_primers).sum()):,}. Percentages are of the annotated "
      "subset, not the corpus (§3.5).")
rm = []
for p, g_ in d.groupby('instrument_platform'):
    v = g_.rc.dropna(); v = v[v > 1]
    rm.append({'Platform': p, 'Runs': len(g_), 'Runs used': len(v),
               'Median Read Count': int(v.median()) if len(v) else None,
               'Q1': int(np.percentile(v,25)) if len(v) else None,
               'Q3': int(np.percentile(v,75)) if len(v) else None})
allv = d.rc.dropna(); allv = allv[allv > 1]
rm.append({'Platform':'ALL PLATFORMS','Runs':N,'Runs used':len(allv),
           'Median Read Count':int(allv.median()),'Q1':int(np.percentile(allv,25)),
           'Q3':int(np.percentile(allv,75))})
sheet(wb, 'run metrics', pd.DataFrame(rm),
      'Medians and quartiles computed on runs with read_count > 1; values of 0 and 1 are archive placeholders.')
tt = pd.DataFrame({'Metric': ['Long-read fraction (% of BioProjects)',
                              "Technology diversity (Shannon H', instrument model)",
                              'Paired-end layout fraction','RANDOM library selection fraction'],
  'End value (2024)': [F12['longread_bp_pct'][-1], F12['shannon_model'][-1],
                       F12['paired_pct'][-1], F12['random_pct'][-1]],
  'MK tau': [F12['mk'][k]['tau'] for k in ['longread','shannon','paired','random']],
  'MK p': [F12['mk'][k]['p'] for k in ['longread','shannon','paired','random']]})
sheet(wb, 'tech trends', tt)
dup = pd.DataFrame(R['duplicates']['top_suspects'])
sheet(wb, 'Duplicate suspects', dup,
      f"{R['duplicates']['n_titles_in_multiple_bps']} titles appear in two or more BioProjects.")
te = pd.DataFrame(F11['residuals']).T.rename_axis('Platform').reset_index()
sheet(wb, 'tech ecosystem', te,
      f"chi2 = {F11['chi2']}, dof = {F11['dof']}, Cramer's V = {F11['V']}, n = {F11['n']:,} runs. "
      "The BGI x Human residual falls from +47.7 in v1.1 to +14.9 here after the 1,067-run SARS-CoV-2 "
      "project was removed; associations attributed to genuine platform preference are unchanged.")
wb.save(f'{W}/Supplementary_Table_S2.xlsx'); print('S2 rebuilt')

# ------------------------------------------------------------------- S3
wb = Workbook(); wb.remove(wb.active)
bt = []
for cat, c in BS['by_category'].items():
    for t in c['types']: bt.append({'broad_category': cat, 'biosample_type': t['biosample_type'],
                                    'n_runs': t['n_runs']})
sheet(wb, 'biosample_types_per_category', pd.DataFrame(bt), NOTE)
sheet(wb, 'country and category', pd.crosstab(d.country, d.broad_category)
      .loc[d.country.value_counts().index].reset_index())
sheet(wb, 'year and platform', pd.crosstab(d.yr, d.instrument_platform).reset_index())
bp = d.bioproject.value_counts()
bins = {'1':0,'2-4':0,'5-9':0,'10-24':0,'25-49':0,'50-99':0,'100-499':0,'500+':0}
for v in bp.values:
    k=('1' if v==1 else '2-4' if v<=4 else '5-9' if v<=9 else '10-24' if v<=24 else '25-49' if v<=49
       else '50-99' if v<=99 else '100-499' if v<=499 else '500+'); bins[k]+=1
sheet(wb, 'bioproject_size_distribution', pd.DataFrame(bins.items(), columns=['index','count']))
top = (d.groupby('bioproject').agg(n_runs=('run_accession','size'), country=('country','first'),
       category=('broad_category','first'), study_title=('study_title','first'))
       .sort_values('n_runs', ascending=False).head(30).reset_index())
sheet(wb, 'top bioprojects', top)
sheet(wb, 'diversity indices', pd.DataFrame(R['diversity']['per_country']))
sheet(wb, 'chi residuals', pd.DataFrame(R['chi_square_residuals']['residuals']),
      f"chi2 = {R['chi_square_residuals']['chi2']:,}, df = {R['chi_square_residuals']['dof']}, "
      f"Cramer's V = {(R['chi_square_residuals']['chi2']/(N*8))**0.5:.3f}")
sheet(wb, 'indicator genome', pd.DataFrame(R['indicator_species']['top']),
      'Reported as descriptive signals; none survives false-discovery correction across the 80-type candidate set.')
sheet(wb, 'mann kendall', pd.DataFrame(R['mann_kendall']))
pc = pd.DataFrame({k: v for k, v in PC.items() if isinstance(v, dict) and 'n_bioprojects' in v}).T
pc.index.name = 'category'; sheet(wb, 'per category permanova', pc.reset_index(),
  'Clinical retains no country with the three or more qualifying BioProjects the test requires once the '
  'misclassified records are removed, so the family comprises six tests. PERMANOVA and PERMDISP use '
  'scikit-bio; BH correction is applied across the six testable categories.')
q = R['permanova']
sheet(wb, 'all_category_permanova', pd.DataFrame([{'Test':'PERMANOVA (999 perm)',
      'n_BioProjects':q['n_bioprojects'],'n_countries':q['n_countries'],
      'F':q['F_statistic'],'R2_pct':round(100*q['R_squared'],2),'p':q['p_value']}]))
wb.save(f'{W}/Supplementary_Table_S3.xlsx'); print('S3 rebuilt')

# ------------------------------------------------------------------- S4 note
wb = load_workbook(f'{W}/Supplementary_Table_S4_BioProject_validation.xlsx')
ws = wb['Summary']; ws.append([]); ws.append(['ROUND 2 UPDATE', ''])
for a, b in [('Action taken','All 4,423 runs in studies judged not to be community metagenomes have been '
                             'REMOVED from the analytical corpus (they were retained in v1.1).'),
             ('Resulting corpus','56,494 runs, 49,308 samples, 1,397 BioProjects'),
             ('Precision','No longer reported: the census is applied as a curation step rather than as an '
                          'error rate, so the released corpus contains only records two independent '
                          'annotators judged to be community metagenomes.'),
             ('See also','Supplementary Table S5 (rejected-arm screen, 791 runs recovered) and S6 '
                         '(broad-category validation, 360 runs).')]:
    ws.append([a, b])
wb.save(f'{W}/Supplementary_Table_S4_BioProject_validation.xlsx'); print('S4 summary updated')
